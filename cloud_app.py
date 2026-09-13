"""
tMart Rebate DN Dashboard — Cloud Edition
==========================================
• Runs on Render.com (free tier) with LibreOffice for PDF conversion
• Each colleague gets their own private session
• Email auto-sent via Gmail OAuth (popup sign-in) or App Password
"""

import os, re, io, json, uuid, shutil, tempfile, zipfile, threading, datetime
import smtplib, urllib.request as urlreq, urllib.parse, base64 as b64lib
from email.mime.multipart import MIMEMultipart
from email.mime.text       import MIMEText
from email.mime.application import MIMEApplication
from flask import Flask, jsonify, send_file, request, session, redirect

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "tmart-dn-cloud-change-this-secret")

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
SESSIONS_DIR = os.path.join(tempfile.gettempdir(), "tmart_sessions")
os.makedirs(SESSIONS_DIR, exist_ok=True)

# Google OAuth config — set these as Render environment variables
GOOGLE_CLIENT_ID     = os.environ.get("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.environ.get("GOOGLE_CLIENT_SECRET", "")
GOOGLE_AUTH_URL      = "https://accounts.google.com/o/oauth2/auth"
GOOGLE_TOKEN_URL     = "https://oauth2.googleapis.com/token"
GMAIL_USERINFO_URL   = "https://www.googleapis.com/oauth2/v1/userinfo"
GMAIL_SEND_URL       = "https://gmail.googleapis.com/gmail/v1/users/me/messages/send"

# ── Session helpers ────────────────────────────────────────────────────────────

def get_sid():
    if "sid" not in session:
        session["sid"] = str(uuid.uuid4())[:12]
    return session["sid"]

def sdir(sid=None):
    d = os.path.join(SESSIONS_DIR, sid or get_sid())
    os.makedirs(d, exist_ok=True)
    return d

def sfile(name, sid=None):   return os.path.join(sdir(sid), name)
def outdir(sid=None):
    d = sfile("output", sid)
    os.makedirs(d, exist_ok=True)
    return d
def excel_path(sid=None):  return sfile("DN_DATA.xlsx", sid)
def template_path(sid=None):
    p = sfile("template.docx", sid)
    if os.path.exists(p): return p
    g = os.path.join(BASE_DIR, "DN_INVOICE _FORMAT.docx")
    return g if os.path.exists(g) else None

# ── Value helpers ──────────────────────────────────────────────────────────────

def str_val(v):
    if v is None: return ""
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime("%d.%m.%Y")
    if isinstance(v, float): return str(int(v)) if v == int(v) else f"{v:.3f}".rstrip("0").rstrip(".")
    if isinstance(v, int): return str(v)
    return str(v).strip()

def col(row, name):
    if name in row: return row[name]
    ns = name.strip()
    if ns in row: return row[ns]
    nl = ns.lower()
    for k, v in row.items():
        if k.strip().lower() == nl: return v
    return ""

def pdf_name(row):
    si = str(row.get("SI No") or row.get("_row_idx", "")).strip()
    fn = col(row, "File Name") or col(row, "Supplier name") or "unknown"
    fn = re.sub(r'[\\/:*?"<>|]', "_", str(fn)).strip()
    return f"{si}_{fn}.pdf"

# ── Excel reading ──────────────────────────────────────────────────────────────

def read_excel(sid=None):
    ep = excel_path(sid)
    if not os.path.exists(ep):
        return {"error": "No Excel uploaded — click 📂 Upload Excel"}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ep, data_only=True)
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
        headers = [str(c.value).strip() if c.value else f"col{i}" for i, c in enumerate(ws[1])]
        rows = []
        for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
            if all(v is None for v in r): continue
            row = {h: str_val(v) for h, v in zip(headers, r)}
            row["_row_idx"] = str(idx)
            rows.append(row)
        wb.close()
        return rows
    except Exception as e:
        return {"error": str(e)}

# ── Template filling ───────────────────────────────────────────────────────────

def fill_template(template, output, row):
    from docx import Document
    total_str = col(row, "Total rebate amount")
    try:    is_credit = float(total_str) < 0
    except: is_credit = False

    reps = {
        "<<BREPLACE>>": col(row, "Invoice Date"),
        "<<CREPLACE>>": col(row, "Finance ID"),
        "<<DREPLACE>>": col(row, "Supplier name"),
        "<<EREPLACE>>": col(row, "Exlusive -0%"),
        "<<FREPLACE>>": col(row, "Exlusive-10%"),
        "<<GREPLACE>>": col(row, "VAT Amount"),
        "<<HREPLACE>>": col(row, "Ex 10%+Vat"),
        "<<IREPLACE>>": col(row, "T_rebate_excl"),
        "<<JREPLACE>>": col(row, "Total rebate amount"),
        "<<KREPLACE>>": col(row, "Debit note number"),
        "<<LREPLACE>>": col(row, "Description"),
        "<<MREPLACE>>": col(row, "VAT No."),
        "<<NREPLACE>>": col(row, "Address"),
    }

    def process(para):
        full = "".join(r.text for r in para.runs)
        if not full.strip(): return
        changed = False
        for old, new in reps.items():
            if old in full:
                full = full.replace(old, str(new) if new else "")
                changed = True
        if is_credit:
            for dn, cn in (("DEBIT NOTE","CREDIT NOTE"),("Debit Note","Credit Note"),("Debit note","Credit Note")):
                if dn in full: full = full.replace(dn, cn); changed = True
        cleaned = re.sub(r"<<[A-Z]+REPLACE>>", "", full)
        if cleaned != full: full = cleaned; changed = True
        if changed and para.runs:
            para.runs[0].text = full
            for r in para.runs[1:]: r.text = ""

    doc = Document(template)
    for p in doc.paragraphs: process(p)
    for tbl in doc.tables:
        for tr in tbl.rows:
            for cell in tr.cells:
                for p in cell.paragraphs: process(p)
    doc.save(output)

# ── PDF conversion ─────────────────────────────────────────────────────────────

def convert_to_pdf(docx_path, pdf_path):
    import subprocess
    lo = shutil.which("libreoffice") or shutil.which("soffice")
    if lo:
        tmp = tempfile.mkdtemp()
        try:
            r = subprocess.run([lo, "--headless", "--convert-to", "pdf",
                                "--outdir", tmp, os.path.abspath(docx_path)],
                               capture_output=True, text=True, timeout=120)
            base = os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
            src  = os.path.join(tmp, base)
            if os.path.exists(src):
                shutil.move(src, pdf_path)
                return True, None
            return False, r.stderr or "LibreOffice: PDF not created"
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    td = docx_path.replace('"', '`"'); op = pdf_path.replace('"', '`"')
    ps = (f'$w=New-Object -ComObject Word.Application;$w.Visible=$false;$w.DisplayAlerts=0;'
          f'$d=$w.Documents.Open("{td}");$d.ExportAsFixedFormat("{op}",17);$d.Close($false);$w.Quit()')
    import subprocess
    r = subprocess.run(["powershell","-ExecutionPolicy","Bypass","-Command",ps],capture_output=True,text=True,timeout=120)
    if os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0: return True, None
    return False, r.stderr or "Word COM: PDF not created"

# ── PDF generation ─────────────────────────────────────────────────────────────

def generate_one(row, sid=None):
    tmpl = template_path(sid)
    if not tmpl:
        return {"pdf_name": pdf_name(row), "ok": False, "error": "No template"}
    pn = pdf_name(row); od = outdir(sid)
    out_pdf = os.path.join(od, pn); temp_doc = os.path.join(od, pn.replace(".pdf","_filled.docx"))
    try:
        fill_template(tmpl, temp_doc, row)
        ok, err = convert_to_pdf(temp_doc, out_pdf)
        if ok and os.path.exists(out_pdf) and os.path.getsize(out_pdf) > 0:
            return {"pdf_name": pn, "ok": True, "error": None}
        return {"pdf_name": pn, "ok": False, "error": err or "PDF empty"}
    except Exception as e:
        return {"pdf_name": pn, "ok": False, "error": str(e)}
    finally:
        try:
            if os.path.exists(temp_doc): os.remove(temp_doc)
        except: pass

_jobs = {}

def _run_batch(sid, rows_list):
    status = _jobs[sid]
    for row in rows_list:
        if status.get("cancelled"): break
        r = generate_one(row, sid=sid)
        status["done"] += 1; status["results"].append(r)
        if not r["ok"]: status["errors"].append(r)
        _write_job(sid, status)
    status["running"] = False; _write_job(sid, status)

def _write_job(sid, status):
    path = os.path.join(SESSIONS_DIR, sid, "job.json")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f: json.dump(status, f)

def _read_job(sid):
    path = os.path.join(SESSIONS_DIR, sid, "job.json")
    try:
        with open(path) as f: return json.load(f)
    except:
        return {"total":0,"done":0,"running":False,"results":[],"errors":[]}

# ── Gmail send helpers ─────────────────────────────────────────────────────────

def _build_mime(to, cc, subject, body_text, pdf_path):
    msg = MIMEMultipart()
    msg["To"] = to
    if cc: msg["CC"] = cc
    msg["Subject"] = subject
    msg.attach(MIMEText(body_text, "plain"))
    if pdf_path and os.path.exists(pdf_path):
        with open(pdf_path, "rb") as f:
            part = MIMEApplication(f.read(), Name=os.path.basename(pdf_path))
        part["Content-Disposition"] = f'attachment; filename="{os.path.basename(pdf_path)}"'
        msg.attach(part)
    return msg

def send_gmail_smtp(to, cc, subject, body_text, pdf_path, gmail_user, gmail_pass):
    msg = _build_mime(to, cc, subject, body_text, pdf_path)
    msg["From"] = gmail_user
    recipients = [a.strip() for a in (to + ";" + cc).split(";") if a.strip()]
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as s:
        s.login(gmail_user, gmail_pass)
        s.sendmail(gmail_user, recipients, msg.as_string())

def send_gmail_api(access_token, to, cc, subject, body_text, pdf_path):
    msg = _build_mime(to, cc, subject, body_text, pdf_path)
    raw = b64lib.urlsafe_b64encode(msg.as_bytes()).decode().rstrip("=")
    payload = json.dumps({"raw": raw}).encode()
    req = urlreq.Request(GMAIL_SEND_URL, data=payload,
                         headers={"Authorization": f"Bearer {access_token}",
                                  "Content-Type": "application/json"})
    try:
        with urlreq.urlopen(req, timeout=30) as resp:
            return True, json.loads(resp.read()).get("id", "sent")
    except Exception as e:
        err = ""
        if hasattr(e, "read"):
            try: err = json.loads(e.read()).get("error", {}).get("message", "")
            except: pass
        return False, err or str(e)

# ── OAuth helpers ──────────────────────────────────────────────────────────────

def _get_redirect_uri():
    base = request.host_url.rstrip("/")
    return base + "/auth/google/callback"

def _get_valid_token():
    """Return a valid OAuth access token from session, refreshing if expired."""
    access_token   = session.get("oauth_access_token", "")
    refresh_token  = session.get("oauth_refresh_token", "")
    expiry_str     = session.get("oauth_expiry", "")
    if not access_token:
        return None
    if expiry_str:
        try:
            expiry = datetime.datetime.fromisoformat(expiry_str)
            if datetime.datetime.now() >= expiry and refresh_token and GOOGLE_CLIENT_SECRET:
                data = urllib.parse.urlencode({
                    "refresh_token": refresh_token,
                    "client_id": GOOGLE_CLIENT_ID,
                    "client_secret": GOOGLE_CLIENT_SECRET,
                    "grant_type": "refresh_token"
                }).encode()
                req = urlreq.Request(GOOGLE_TOKEN_URL, data=data,
                                     headers={"Content-Type": "application/x-www-form-urlencoded"})
                with urlreq.urlopen(req, timeout=30) as resp:
                    tok = json.loads(resp.read())
                access_token = tok["access_token"]
                session["oauth_access_token"] = access_token
                session["oauth_expiry"] = (datetime.datetime.now() +
                    datetime.timedelta(seconds=tok.get("expires_in", 3600) - 60)).isoformat()
        except: pass
    return access_token

# ── API: OAuth flow ────────────────────────────────────────────────────────────

@app.route("/auth/google")
def auth_google():
    if not GOOGLE_CLIENT_ID:
        return ("<html><body style='font-family:sans-serif;padding:30px'>"
                "<h2 style='color:#E05500'>⚠ Google OAuth Not Configured</h2>"
                "<p>Ask IT to add <b>GOOGLE_CLIENT_ID</b> and <b>GOOGLE_CLIENT_SECRET</b> "
                "as environment variables on Render.</p>"
                "<p>The Redirect URI to register in Google Cloud Console is:<br>"
                f"<code style='background:#eee;padding:4px 8px'>{request.host_url.rstrip('/')}/auth/google/callback</code></p>"
                "</body></html>"), 500
    params = urllib.parse.urlencode({
        "client_id":     GOOGLE_CLIENT_ID,
        "redirect_uri":  _get_redirect_uri(),
        "scope":         "https://www.googleapis.com/auth/gmail.send https://www.googleapis.com/auth/userinfo.email",
        "response_type": "code",
        "access_type":   "offline",
        "prompt":        "consent",
        "state":         get_sid()
    })
    return redirect(GOOGLE_AUTH_URL + "?" + params)

@app.route("/auth/google/callback")
def auth_google_callback():
    code  = request.args.get("code", "")
    error = request.args.get("error", "")
    CLOSE = "<html><body><script>window.opener&&window.opener.postMessage({t},{o});setTimeout(()=>window.close(),400);</script><p>{m}</p></body></html>"

    if error:
        return CLOSE.format(t=json.dumps({"ok":False,"error":error}), o="'*'", m=f"Error: {error}")
    if not code:
        return CLOSE.format(t=json.dumps({"ok":False,"error":"No code"}), o="'*'", m="No code returned")

    try:
        # Exchange code for tokens
        data = urllib.parse.urlencode({
            "code": code, "client_id": GOOGLE_CLIENT_ID,
            "client_secret": GOOGLE_CLIENT_SECRET,
            "redirect_uri": _get_redirect_uri(), "grant_type": "authorization_code"
        }).encode()
        req = urlreq.Request(GOOGLE_TOKEN_URL, data=data,
                             headers={"Content-Type": "application/x-www-form-urlencoded"})
        with urlreq.urlopen(req, timeout=30) as resp:
            tok = json.loads(resp.read())

        access_token  = tok["access_token"]
        refresh_token = tok.get("refresh_token", session.get("oauth_refresh_token", ""))

        # Get user email
        req2 = urlreq.Request(GMAIL_USERINFO_URL, headers={"Authorization": f"Bearer {access_token}"})
        with urlreq.urlopen(req2, timeout=10) as resp2:
            info = json.loads(resp2.read())
        email = info.get("email", "")

        # Store in session
        session["gmail_mode"]          = "oauth"
        session["gmail_user"]          = email
        session["gmail_pass"]          = ""
        session["oauth_access_token"]  = access_token
        session["oauth_refresh_token"] = refresh_token
        session["oauth_expiry"]        = (datetime.datetime.now() +
            datetime.timedelta(seconds=tok.get("expires_in", 3600) - 60)).isoformat()

        msg = json.dumps({"ok": True, "email": email})
        return CLOSE.format(t=msg, o="'*'",
                            m=f"✅ Connected as {email}. Closing…")
    except Exception as e:
        err = str(e)
        return CLOSE.format(t=json.dumps({"ok":False,"error":err}), o="'*'", m=f"Error: {err}")

# ── API routes ─────────────────────────────────────────────────────────────────

@app.route("/")
def index(): return DASHBOARD_HTML

@app.route("/api/session-id")
def api_session_id(): return jsonify({"sid": get_sid()})

@app.route("/api/rows")
def api_rows():
    rows = read_excel()
    if isinstance(rows, dict): return jsonify({"ok": False, "error": rows["error"]})
    od = outdir(); sent = session.get("sent", {})
    result = []
    for i, row in enumerate(rows):
        pn = pdf_name(row)
        result.append({
            "idx": i, "si_no": row.get("_row_idx", str(i+1)),
            "supplier": col(row, "Supplier name"), "dn_number": col(row, "Debit note number"),
            "email": col(row, "Email") or col(row, "email"),
            "cc": col(row, "cc") or col(row, "CC") or col(row, "Cc") or "",
            "pdf_name": pn, "pdf_exists": os.path.exists(os.path.join(od, pn)),
            "email_sent": sent.get(pn, False), "row": row,
        })
    return jsonify({"ok": True, "rows": result})

@app.route("/api/generate-start", methods=["POST"])
def api_generate_start():
    body = request.json or {}; rows_list = [item["row"] for item in body.get("rows", [])]
    sid  = get_sid()
    status = {"total": len(rows_list), "done": 0, "running": True, "results": [], "errors": [], "cancelled": False}
    _jobs[sid] = status; _write_job(sid, status)
    threading.Thread(target=_run_batch, args=(sid, rows_list), daemon=True).start()
    return jsonify({"ok": True, "total": len(rows_list)})

@app.route("/api/generate-status")
def api_generate_status(): return jsonify(_read_job(get_sid()))

@app.route("/api/generate", methods=["POST"])
def api_generate():
    body = request.json or {}; rows_list = [item["row"] for item in body.get("rows", [])]
    sid  = get_sid()
    return jsonify({"results": [generate_one(r, sid=sid) for r in rows_list]})

@app.route("/api/send-email", methods=["POST"])
def api_send_email():
    body = request.json or {}
    pn   = body.get("pdf_name", "")
    pdf_path = os.path.join(outdir(), pn)
    if not os.path.exists(pdf_path):
        return jsonify({"ok": False, "error": "PDF not found — generate it first"})

    to       = body.get("email", "")
    cc       = body.get("cc", "")
    supplier = body.get("supplier", "Supplier")
    dn_num   = body.get("dn_number", "")
    subject  = f"Rebate Debit Note {dn_num} - DH Store Bahrain (tMart)".strip()
    body_t   = (f"Dear {supplier},\n\nPlease find attached your Rebate Debit Note.\n\nRegards,\ntMart Finance Team")

    mode = session.get("gmail_mode", "smtp")

    if mode == "oauth":
        token = _get_valid_token()
        if not token:
            return jsonify({"ok": False, "error": "Gmail not connected — click Connect Gmail"})
        ok, result = send_gmail_api(token, to, cc, subject, body_t, pdf_path)
    elif mode == "smtp":
        gu = session.get("gmail_user", ""); gp = session.get("gmail_pass", "")
        if not gu or not gp:
            return jsonify({"ok": False, "error": "Gmail not configured — click ⚙ Gmail Settings"})
        try:
            send_gmail_smtp(to, cc, subject, body_t, pdf_path, gu, gp)
            ok, result = True, "sent"
        except Exception as e:
            ok, result = False, str(e)
    else:
        return jsonify({"ok": False, "error": "Use browser mode from the frontend"})

    if ok:
        sent = session.get("sent", {}); sent[pn] = True; session["sent"] = sent
        return jsonify({"ok": True})
    return jsonify({"ok": False, "error": result})

@app.route("/api/gmail-settings", methods=["POST"])
def api_gmail_settings():
    body = request.json or {}; mode = body.get("mode", "smtp")
    session["gmail_mode"] = mode
    session["gmail_user"] = body.get("user", "").strip()
    session["gmail_pass"] = body.get("password", "").strip() if mode == "smtp" else ""
    return jsonify({"ok": True, "user": session["gmail_user"], "mode": mode})

@app.route("/api/gmail-status")
def api_gmail_status():
    mode = session.get("gmail_mode", "")
    user = session.get("gmail_user", "")
    has_oauth = bool(session.get("oauth_access_token"))
    oauth_configured = bool(GOOGLE_CLIENT_ID)
    if mode == "oauth": configured = has_oauth
    elif mode == "smtp": configured = bool(user and session.get("gmail_pass"))
    else: configured = bool(user)
    return jsonify({"configured": configured, "user": user, "mode": mode,
                    "has_oauth": has_oauth, "oauth_configured": oauth_configured})

@app.route("/api/upload-excel", methods=["POST"])
def api_upload_excel():
    try:
        import openpyxl
        f = request.files.get("file")
        if not f: return jsonify({"ok": False, "error": "No file"})
        dest = excel_path(); f.save(dest)
        wb = openpyxl.load_workbook(dest, data_only=True)
        sheets = wb.sheetnames; wb.close()
        if "Sheet1" not in sheets:
            return jsonify({"ok": False, "error": f"Sheet1 not found. Available: {sheets}"})
        return jsonify({"ok": True})
    except Exception as e: return jsonify({"ok": False, "error": str(e)})

@app.route("/api/upload-template", methods=["POST"])
def api_upload_template():
    f = request.files.get("file")
    if not f: return jsonify({"ok": False, "error": "No file"})
    f.save(sfile("template.docx")); return jsonify({"ok": True})

@app.route("/api/template-status")
def api_template_status():
    t = template_path()
    return jsonify({"exists": bool(t), "path": os.path.basename(t) if t else ""})

@app.route("/preview/<sid>/<path:filename>")
def preview(sid, filename):
    path = os.path.join(SESSIONS_DIR, sid, "output", filename)
    if os.path.exists(path): return send_file(path, mimetype="application/pdf")
    return jsonify({"error": "Not found"}), 404

@app.route("/download/<sid>/<path:filename>")
def download(sid, filename):
    path = os.path.join(SESSIONS_DIR, sid, "output", filename)
    if os.path.exists(path): return send_file(path, as_attachment=True, download_name=filename)
    return jsonify({"error": "Not found"}), 404

@app.route("/download-all/<sid>")
def download_all(sid):
    od = os.path.join(SESSIONS_DIR, sid, "output")
    pdfs = [f for f in os.listdir(od) if f.endswith(".pdf")] if os.path.exists(od) else []
    if not pdfs: return jsonify({"error": "No PDFs yet"}), 404
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in pdfs: zf.write(os.path.join(od, f), f)
    buf.seek(0)
    return send_file(buf, mimetype="application/zip", as_attachment=True, download_name="Rebate_DNs.zip")

@app.route("/api/delete/<path:filename>", methods=["DELETE"])
def delete_pdf(filename):
    path = os.path.join(outdir(), filename)
    try:
        if os.path.exists(path): os.remove(path)
        return jsonify({"ok": True})
    except Exception as e: return jsonify({"ok": False, "error": str(e)})

# ── Dashboard HTML ─────────────────────────────────────────────────────────────

DASHBOARD_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8"/>
<meta name="viewport" content="width=device-width,initial-scale=1"/>
<title>tMart – Rebate DN Dashboard</title>
<style>
:root{--or:#E05500;--dark:#F0F2F5;--card:#FFFFFF;--bdr:#D8DBE8;--grn:#16A34A;--red:#DC2626;--blu:#2563EB;--pur:#7C3AED;--tl:#0D9488;}
*{box-sizing:border-box;margin:0;padding:0;}
body{font-family:"Segoe UI",sans-serif;background:var(--dark);color:#1A1A2E;min-height:100vh;}
nav{background:var(--card);border-bottom:3px solid var(--or);padding:14px 28px;display:flex;align-items:center;gap:14px;flex-wrap:wrap;}
.sub{font-size:13px;color:#666;flex:1;}
.cards{display:flex;gap:16px;padding:22px 28px 0;flex-wrap:wrap;}
.card{flex:1;min-width:140px;background:var(--card);border:1px solid var(--bdr);border-radius:10px;padding:16px 20px;}
.card .num{font-size:28px;font-weight:700;color:var(--or);}.card .lbl{font-size:11px;color:#666;margin-top:2px;text-transform:uppercase;letter-spacing:.5px;}
.toolbar{display:flex;gap:8px;padding:16px 28px;flex-wrap:wrap;align-items:center;}
input[type=file]{display:none;}
button,label.btn{cursor:pointer;border:none;border-radius:7px;padding:9px 16px;font-size:13px;font-weight:600;transition:opacity .15s;display:inline-flex;align-items:center;gap:5px;}
button:hover,label.btn:hover{opacity:.85;}
.or{background:var(--or);color:#fff;}.grn{background:var(--grn);color:#fff;}.blu{background:var(--blu);color:#fff;}
.pur{background:var(--pur);color:#fff;}.red{background:var(--red);color:#fff;}.tl{background:var(--tl);color:#fff;}
.out{background:transparent;border:1px solid var(--bdr);color:#555;}
.sm{padding:5px 10px;font-size:12px;border-radius:5px;white-space:nowrap;}.sep{flex:1;}
.tbl-wrap{padding:0 28px 40px;overflow-x:auto;}
table{width:100%;border-collapse:collapse;font-size:13px;}
thead th{background:var(--card);color:var(--or);padding:10px 12px;text-align:left;border-bottom:2px solid var(--bdr);white-space:nowrap;}
tbody tr{border-bottom:1px solid var(--bdr);}tbody tr:hover{background:rgba(255,106,0,.05);}
td{padding:9px 12px;vertical-align:middle;}
.acts{white-space:nowrap;display:flex;gap:5px;flex-wrap:wrap;}
.badge{display:inline-block;padding:3px 9px;border-radius:12px;font-size:11px;font-weight:600;}
.b-ok{background:rgba(34,197,94,.15);color:var(--grn);}.b-no{background:rgba(239,68,68,.1);color:var(--red);}.b-wait{background:rgba(255,106,0,.12);color:var(--or);}
.spin{display:inline-block;width:13px;height:13px;border:2px solid rgba(255,255,255,.3);border-top-color:#fff;border-radius:50%;animation:spin .6s linear infinite;}
@keyframes spin{to{transform:rotate(360deg);}}
#toast-wrap{position:fixed;bottom:22px;right:22px;display:flex;flex-direction:column;gap:8px;z-index:999;}
.toast{background:var(--card);border-left:4px solid var(--grn);padding:11px 16px;border-radius:7px;font-size:13px;box-shadow:0 4px 16px rgba(0,0,0,.15);animation:fi .2s;min-width:260px;max-width:420px;word-break:break-word;color:#1A1A2E;}
.toast.err{border-color:var(--red);}@keyframes fi{from{opacity:0;transform:translateY(8px)}to{opacity:1}}
#pw{display:none;min-width:220px;}.pbar{height:6px;background:var(--bdr);border-radius:3px;overflow:hidden;margin-top:6px;}
.pbar .fill{height:100%;background:var(--or);transition:width .4s;}
#modal{display:none;position:fixed;inset:0;background:rgba(0,0,0,.7);z-index:1000;align-items:center;justify-content:center;}
#modal.show{display:flex;}
#modal-inner{background:var(--card);border-radius:12px;width:90%;max-width:900px;height:82vh;display:flex;flex-direction:column;overflow:hidden;}
#modal-hdr{padding:14px 18px;display:flex;justify-content:space-between;align-items:center;border-bottom:1px solid var(--bdr);}
#modal-title{font-weight:700;color:var(--or);}#modal-close{cursor:pointer;font-size:20px;color:#888;background:none;border:none;}
#modal-frame{flex:1;border:none;}
#gmodal{display:none;position:fixed;inset:0;background:rgba(0,0,0,.75);z-index:1001;align-items:center;justify-content:center;}
#gmodal.show{display:flex;}
#gmodal-inner{background:var(--card);border-radius:14px;width:460px;max-height:90vh;overflow-y:auto;padding:28px;box-shadow:0 8px 40px rgba(0,0,0,.4);}
#gmodal h3{color:var(--or);margin-bottom:16px;font-size:17px;}
.tabs{display:flex;gap:6px;margin-bottom:18px;}
.tab-btn{flex:1;padding:10px 6px;border-radius:8px;border:2px solid var(--bdr);background:transparent;color:#888;cursor:pointer;font-size:12px;font-weight:600;text-align:center;line-height:1.4;transition:all .15s;}
.tab-btn.active{border-color:var(--tl);background:var(--tl);color:#fff;}
.panel{display:none;}.panel.show{display:block;}
.field{margin-bottom:14px;}
.field label{display:block;font-size:12px;color:#666;margin-bottom:5px;text-transform:uppercase;letter-spacing:.4px;}
.field input{width:100%;background:#F7F8FA;border:1px solid var(--bdr);border-radius:7px;padding:9px 12px;color:#1A1A2E;font-size:14px;outline:none;}
.field input:focus{border-color:var(--or);}
.hint{font-size:11px;color:#666;margin-top:5px;line-height:1.6;}
.hint a{color:var(--tl);text-decoration:none;}
.connected-box{background:rgba(34,197,94,.1);border:1px solid var(--grn);border-radius:8px;padding:12px 16px;font-size:13px;color:var(--grn);margin-bottom:14px;}
.google-btn{width:100%;padding:12px;background:#4285F4;color:#fff;border:none;border-radius:8px;font-size:14px;font-weight:600;cursor:pointer;display:flex;align-items:center;justify-content:center;gap:10px;margin:10px 0;}
.google-btn:hover{background:#3367D6;}
.google-btn svg{width:20px;height:20px;}
</style>
</head>
<body>
<div style="background:#fff;border-bottom:3px solid #FF6200;padding:8px 0;text-align:center;">
  <img src="data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAVwAAAGBCAYAAAAqtgndAAAAAXNSR0IArs4c6QAAAARnQU1BAACxjwv8YQUAAAAJcEhZcwAADsMAAA7DAcdvqGQAACyWSURBVHhe7d17cFXV/ffx9z7XXE8SyBmVu1xEoYog/OyjzoDWC1LUqUy9AY8dsWO1M2qrM/XSjtqrTq3l59S/OjBtFceieClWsYwVHe0jIuAVqRgVSLjlBjlJzv3s54/EoMuzJUDOOofk85pxxvPde2etITuf2Vlrryxnf0eXi4iIFJzPLIiISGEocEVELFHgiohYosAVEbFEgSsiYokCV0TEEgWuiIglClwREUsUuCIilihwRUQsUeCKiFiiwBURsUSBKyJiiQJXRMQSBa6IiCUKXBERSxS4IiKWKHBFRCxR4IqIWKLAFRGxRIErImKJAldAxBIFroiIJQpcERFLFLgiIpYocEVELFHgiohYosAVEbFEgSsiYokCV0TEEgWuiIglClwREUsUuCIilihwRUQsUeCKiFiiwBURsUSBKyJiiQJXRMQSBa6IiCUKXBERSxS4IiKWKHBFRCxR4IqIWKLAFRGxRIErImKJAldA" alt="talabat mart" style="height:52px;object-fit:contain;"/>
</div>
<nav>
  <span class="sub" style="font-weight:600;">Rebate Debit Note Dashboard</span>
  <span id="gstatus-nav" style="font-size:12px;color:#888;">⚙ Gmail not connected</span>
  <button class="tl" onclick="openGmail()">⚙ Gmail Settings</button>
</nav>
<div class="cards">
  <div class="card"><div class="num" id="c-total">—</div><div class="lbl">Total Rows</div></div>
  <div class="card"><div class="num" id="c-pdf">—</div><div class="lbl">PDFs Generated</div></div>
  <div class="card"><div class="num" id="c-email">—</div><div class="lbl">Emails Sent</div></div>
</div>
<div class="toolbar">
  <label class="btn or">📂 Upload Excel<input type="file" accept=".xlsx" onchange="uploadExcel(this)"/></label>
  <label class="btn out">📄 Upload Template<input type="file" accept=".docx" onchange="uploadTemplate(this)"/></label>
  <button class="or" onclick="generateAll()">⚡ Generate All PDFs</button>
  <button class="grn" onclick="sendAll()">📧 Send All Emails</button>
  <button class="blu" onclick="dlAll()">⬇ Download All (ZIP)</button>
  <div class="sep"></div>
  <div id="pw">
    <div style="font-size:12px;color:#666;" id="pl">Working…</div>
    <div class="pbar"><div class="fill" id="pb" style="width:0%"></div></div>
  </div>
  <button class="out" onclick="loadRows()">↺ Refresh</button>
</div>
<div class="tbl-wrap">
<table>
  <thead>
    <tr>
      <th><input type="checkbox" id="chk-all" onchange="toggleAll(this)"></th>
      <th>#</th><th>Supplier</th><th>DN Number</th><th>Email</th>
      <th>PDF</th><th>Email Sent</th><th>Actions</th>
    </tr>
  </thead>
  <tbody id="tbody">
    <tr><td colspan="8" style="text-align:center;padding:40px;color:#666;">Upload DN_DATA.xlsx to begin</td></tr>
  </tbody>
</table>
</div>

<!-- PDF Preview Modal -->
<div id="modal"><div id="modal-inner">
  <div id="modal-hdr">
    <span id="modal-title">PDF Preview</span>
    <button id="modal-close" onclick="closeModal()">✕</button>
  </div>
  <iframe id="modal-frame" src=""></iframe>
</div></div>

<!-- Gmail Settings Modal -->
<div id="gmodal"><div id="gmodal-inner">
  <h3>⚙ Gmail Settings</h3>
  <div class="tabs">
    <button class="tab-btn active" id="tab-oauth" onclick="setTab('oauth')">🔗 Connect Google<br><span style="font-weight:400;font-size:10px;">Popup sign-in</span></button>
    <button class="tab-btn" id="tab-smtp" onclick="setTab('smtp')">🔑 App Password<br><span style="font-weight:400;font-size:10px;">SMTP auto-send</span></button>
    <button class="tab-btn" id="tab-browser" onclick="setTab('browser')">🌐 Browser<br><span style="font-weight:400;font-size:10px;">Manual attach</span></button>
  </div>

  <!-- OAuth panel -->
  <div id="panel-oauth" class="panel show">
    <div id="oauth-connected-box" class="connected-box" style="display:none;">
      ✅ Connected as <strong id="oauth-email-disp"></strong><br>
      <span style="font-size:12px;color:#555;">Emails will send automatically with PDF attached.</span>
    </div>
    <div id="oauth-action-box">
      <p style="font-size:13px;color:#444;margin-bottom:12px;line-height:1.6;">
        Click below to sign in with your Google account in a popup window.<br>
        Once you approve access, all emails will send automatically with the PDF attached — no manual steps.
      </p>
      <button class="google-btn" onclick="connectGooglePopup()">
        <svg viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg"><path fill="#fff" d="M22.56 12.25c0-.78-.07-1.53-.2-2.25H12v4.26h5.92c-.26 1.37-1.04 2.53-2.21 3.31v2.77h3.57c2.08-1.92 3.28-4.74 3.28-8.09z"/><path fill="#fff" d="M12 23c2.97 0 5.46-.98 7.28-2.66l-3.57-2.77c-.98.66-2.23 1.06-3.71 1.06-2.86 0-5.29-1.93-6.16-4.53H2.18v2.84C3.99 20.53 7.7 23 12 23z"/><path fill="#fff" d="M5.84 14.09c-.22-.66-.35-1.36-.35-2.09s.13-1.43.35-2.09V7.07H2.18C1.43 8.55 1 10.22 1 12s.43 3.45 1.18 4.93l3.66-2.84z"/><path fill="#fff" d="M12 5.38c1.62 0 3.06.56 4.21 1.64l3.15-3.15C17.45 2.09 14.97 1 12 1 7.7 1 3.99 3.47 2.18 7.07l3.66 2.84c.87-2.6 3.3-4.53 6.16-4.53z"/></svg>
        Sign in with Google
      </button>
      <div id="oauth-not-configured" style="display:none;background:rgba(239,68,68,.08);border:1px solid var(--red);border-radius:8px;padding:12px;font-size:12px;color:#7f1d1d;line-height:1.7;">
        ⚠️ <strong>Google OAuth not configured on the server.</strong><br>
        Ask IT to add <code>GOOGLE_CLIENT_ID</code> and <code>GOOGLE_CLIENT_SECRET</code> as Render environment variables,<br>
        with this Redirect URI registered in Google Cloud Console:<br>
        <code id="redirect-uri-display" style="word-break:break-all;"></code>
      </div>
    </div>
    <div style="display:flex;gap:10px;margin-top:18px;">
      <button class="out" onclick="closeGmail()" style="flex:1;">Close</button>
    </div>
  </div>

  <!-- SMTP panel -->
  <div id="panel-smtp" class="panel">
    <div class="field"><label>Gmail Address</label><input type="email" id="g-user-smtp" placeholder="yourname@gmail.com"/></div>
    <div class="field">
      <label>Gmail App Password</label>
      <input type="password" id="g-pass" placeholder="xxxx xxxx xxxx xxxx"/>
      <div class="hint">Get it at: <a href="https://myaccount.google.com/apppasswords" target="_blank">myaccount.google.com/apppasswords</a><br>(Requires 2-Step Verification on your Google account)</div>
    </div>
    <div style="display:flex;gap:10px;margin-top:14px;">
      <button class="tl" onclick="saveSmtp()" style="flex:1;">💾 Save &amp; Connect</button>
      <button class="out" onclick="closeGmail()" style="flex:1;">Cancel</button>
    </div>
  </div>

  <!-- Browser panel -->
  <div id="panel-browser" class="panel">
    <div class="field"><label>Your Gmail Address</label><input type="email" id="g-user-browser" placeholder="yourname@gmail.com"/></div>
    <div class="hint" style="margin-top:8px;">ℹ️ Manual mode — Gmail opens pre-filled in a new tab. Attach the PDF and click Send yourself.<br><br>For fully automatic sending, use <strong>Connect Google</strong> or <strong>App Password</strong>.</div>
    <div style="display:flex;gap:10px;margin-top:18px;">
      <button class="tl" onclick="saveBrowser()" style="flex:1;">💾 Save</button>
      <button class="out" onclick="closeGmail()" style="flex:1;">Cancel</button>
    </div>
  </div>
</div></div>

<div id="toast-wrap"></div>
<script>
let rows=[], SID='', _gmailMode='oauth', _oauthConnected=false;

(async()=>{
  const r = await fetch('/api/session-id').then(r=>r.json());
  SID = r.sid;
  loadRows();
  checkGmail();
  // Listen for OAuth popup message
  window.addEventListener('message', e => {
    if(e.data && e.data.ok !== undefined){
      if(e.data.ok){
        _oauthConnected = true;
        showOAuthConnected(e.data.email);
        checkGmail();
        toast('✅ Connected as '+e.data.email+' — ready to send!');
      } else {
        toast('Google sign-in failed: '+(e.data.error||'unknown'), true);
      }
    }
  });
})();

async function loadRows(){
  const r = await fetch('/api/rows').then(r=>r.json());
  if(!r.ok){toast('Error: '+r.error,true);return;}
  rows=r.rows; render(); cards();
}
function cards(){
  document.getElementById('c-total').textContent=rows.length;
  document.getElementById('c-pdf').textContent=rows.filter(r=>r.pdf_exists).length;
  document.getElementById('c-email').textContent=rows.filter(r=>r.email_sent).length;
}
function render(){
  const b=document.getElementById('tbody');
  if(!rows.length){b.innerHTML='<tr><td colspan="8" style="text-align:center;padding:40px;color:#666;">No data — upload DN_DATA.xlsx</td></tr>';return;}
  b.innerHTML=rows.map((r,i)=>`
  <tr id="tr-${i}">
    <td><input type="checkbox" class="rc" data-idx="${i}"></td>
    <td style="color:#666">${r.si_no}</td>
    <td><strong>${r.supplier}</strong></td>
    <td style="color:#888;font-size:12px">${r.dn_number||''}</td>
    <td style="font-size:11px;color:#555;max-width:240px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap" title="${r.email}">${r.email}</td>
    <td>${r.pdf_exists?'<span class="badge b-ok">✓ Ready</span>':'<span class="badge b-no">Pending</span>'}</td>
    <td>${r.email_sent?'<span class="badge b-ok">✓ Sent</span>':'<span class="badge b-wait">Not sent</span>'}</td>
    <td class="acts">
      <button class="or sm" id="gb-${i}" onclick="genOne(${i})">⚡ Generate</button>
      ${r.pdf_exists?`<button class="pur sm" onclick="previewOne('${r.pdf_name}','${r.supplier}')">👁 Preview</button>`:''}
      ${r.pdf_exists?`<button class="blu sm" onclick="dlOne('${r.pdf_name}')">⬇ Download</button>`:''}
      ${r.pdf_exists?`<button class="grn sm" id="mb-${i}" onclick="sendOne(${i})">📧 Send</button>`:''}
      ${r.pdf_exists?`<button class="red sm" onclick="delOne(${i},'${r.pdf_name}')">🗑 Delete</button>`:''}
    </td>
  </tr>`).join('');
}

function sel(){
  const idx=[...document.querySelectorAll('.rc:checked')].map(c=>+c.dataset.idx);
  return idx.length?idx.map(i=>rows[i]):rows;
}
function toggleAll(cb){document.querySelectorAll('.rc').forEach(c=>c.checked=cb.checked);}

async function genOne(i){
  const btn=document.getElementById(`gb-${i}`);
  btn.innerHTML='<span class="spin"></span>';btn.disabled=true;
  try{
    const res=await post('/api/generate',{rows:[rows[i]]});
    const r0=res.results[0];
    if(r0.ok){rows[i].pdf_exists=true;toast('✓ PDF ready: '+r0.pdf_name);}
    else toast('Error: '+r0.error,true);
  }catch(e){toast('Error: '+e.message,true);}
  btn.innerHTML='⚡ Generate';btn.disabled=false;render();cards();
}

async function generateAll(){
  const s=sel();if(!s.length){toast('No rows',true);return;}
  prog(0,s.length,'Starting…');
  try{
    const res=await post('/api/generate-start',{rows:s});
    if(!res.ok){hideProg();toast('Error starting',true);return;}
    const total=res.total;
    const poll=setInterval(async()=>{
      const st=await fetch('/api/generate-status').then(r=>r.json());
      prog(st.done,total,`Generating… ${st.done}/${total}`);
      if(!st.running){
        clearInterval(poll);
        st.results.forEach(r=>{if(r.ok){const m=rows.find(x=>x.pdf_name===r.pdf_name);if(m)m.pdf_exists=true;}});
        const ok=st.results.filter(r=>r.ok).length,err=st.results.filter(r=>!r.ok).length;
        hideProg();
        toast(err===0?`✓ ${ok} PDFs generated`:`${ok} ok, ${err} failed`,err>0);
        if(err>0)st.errors.slice(0,3).forEach(r=>toast(r.pdf_name+': '+r.error,true));
        render();cards();
      }
    },2500);
  }catch(e){hideProg();toast('Error: '+e.message,true);}
}

function previewOne(pn,sup){
  document.getElementById('modal-title').textContent=sup+' — '+pn;
  document.getElementById('modal-frame').src=`/preview/${SID}/${encodeURIComponent(pn)}`;
  document.getElementById('modal').classList.add('show');
}
function closeModal(){document.getElementById('modal').classList.remove('show');document.getElementById('modal-frame').src='';}
document.getElementById('modal').addEventListener('click',function(e){if(e.target===this)closeModal();});

function dlOne(pn){window.open(`/download/${SID}/${encodeURIComponent(pn)}`,'_blank');}
function dlAll(){window.open(`/download-all/${SID}`,'_blank');}

async function delOne(i,pn){
  if(!confirm('Delete '+pn+'?'))return;
  await fetch(`/api/delete/${encodeURIComponent(pn)}`,{method:'DELETE'});
  rows[i].pdf_exists=false;rows[i].email_sent=false;render();cards();toast('🗑 Deleted');
}

function gmailComposeUrl(to,cc,subject,body){
  const p=new URLSearchParams({view:'cm',to,cc,su:subject,body});
  return 'https://mail.google.com/mail/?'+p.toString();
}

async function sendOne(i){
  const row=rows[i];
  if(!row.pdf_exists){toast('Generate PDF first',true);return;}
  const btn=document.getElementById(`mb-${i}`);
  btn.innerHTML='<span class="spin"></span>';btn.disabled=true;
  const subject=`Rebate Debit Note ${row.dn_number||''} - DH Store Bahrain (tMart)`.trim();
  const bodyTxt=`Dear ${row.supplier},\\n\\nPlease find attached your Rebate Debit Note.\\n\\nRegards,\\ntMart Finance Team`;
  try{
    if(_gmailMode==='browser'){
      window.open(`/download/${SID}/${encodeURIComponent(row.pdf_name)}`,'_blank');
      setTimeout(()=>window.open(gmailComposeUrl(row.email,row.cc,subject,bodyTxt),'_blank'),800);
      rows[i].email_sent=true;
      toast('📥 PDF downloaded — attach it in the Gmail tab that opened');
    } else {
      const res=await post('/api/send-email',{pdf_name:row.pdf_name,email:row.email,cc:row.cc,supplier:row.supplier,dn_number:row.dn_number});
      if(res.ok){rows[i].email_sent=true;toast('✅ Email sent to '+row.email);}
      else toast('Email failed: '+res.error,true);
    }
  }catch(e){toast('Error: '+e.message,true);}
  btn.innerHTML='📧 Send';btn.disabled=false;render();cards();
}

async function sendAll(){
  const s=sel().filter(r=>r.pdf_exists);
  if(!s.length){toast('No PDFs ready — generate first',true);return;}
  prog(0,s.length,'Sending emails…');
  let sent=0,failed=0;
  for(const row of s){
    try{
      if(_gmailMode==='browser'){
        const subject=`Rebate Debit Note ${row.dn_number||''} - DH Store Bahrain (tMart)`.trim();
        const bodyTxt=`Dear ${row.supplier},\\n\\nPlease find attached your Rebate Debit Note.\\n\\nRegards,\\ntMart Finance Team`;
        window.open(gmailComposeUrl(row.email,row.cc,subject,bodyTxt),'_blank');
        row.email_sent=true;sent++;
        await new Promise(r=>setTimeout(r,600));
      } else {
        const res=await post('/api/send-email',{pdf_name:row.pdf_name,email:row.email,cc:row.cc,supplier:row.supplier,dn_number:row.dn_number});
        if(res.ok){row.email_sent=true;sent++;}
        else{failed++;toast('Failed ('+row.supplier+'): '+res.error,true);}
      }
    }catch(e){failed++;toast('Error: '+e.message,true);}
    prog(sent+failed,s.length,`Sending… ${sent+failed}/${s.length}`);
    await new Promise(r=>setTimeout(r,300));
  }
  hideProg();
  if(_gmailMode==='browser') toast(`Download the ZIP and attach PDFs to the ${sent} Gmail tabs opened`);
  else toast(failed===0?`✅ ${sent} emails sent!`:`${sent} sent, ${failed} failed`,failed>0);
  render();cards();
}

/* ── Gmail OAuth popup ── */
function connectGooglePopup(){
  const w=500,h=600,l=screen.width/2-w/2,t=screen.height/2-h/2;
  window.open('/auth/google','gmailOAuth',`width=${w},height=${h},left=${l},top=${t},menubar=no,toolbar=no,location=no`);
  toast('🔗 Google sign-in window opened — approve access there');
}

function showOAuthConnected(email){
  document.getElementById('oauth-email-disp').textContent=email;
  document.getElementById('oauth-connected-box').style.display='block';
  document.getElementById('oauth-action-box').querySelector('.google-btn').textContent='🔄 Re-connect Google';
}

/* ── Gmail Settings modal ── */
function openGmail(){document.getElementById('gmodal').classList.add('show');}
function closeGmail(){document.getElementById('gmodal').classList.remove('show');}
document.getElementById('gmodal').addEventListener('click',function(e){if(e.target===this)closeGmail();});

function setTab(t){
  ['oauth','smtp','browser'].forEach(x=>{
    document.getElementById('tab-'+x).classList.toggle('active',x===t);
    document.getElementById('panel-'+x).classList.toggle('show',x===t);
  });
  _gmailMode=t;
}

async function saveSmtp(){
  const user=document.getElementById('g-user-smtp').value.trim();
  const pass=document.getElementById('g-pass').value.trim();
  if(!user||!pass){toast('Enter Gmail address and App Password',true);return;}
  await post('/api/gmail-settings',{mode:'smtp',user,password:pass});
  closeGmail();checkGmail();toast('✅ Gmail SMTP connected');
}
async function saveBrowser(){
  const user=document.getElementById('g-user-browser').value.trim();
  if(!user){toast('Enter your Gmail address',true);return;}
  await post('/api/gmail-settings',{mode:'browser',user,password:''});
  closeGmail();checkGmail();toast('✅ Browser Gmail mode saved');
}

async function checkGmail(){
  const r=await fetch('/api/gmail-status').then(r=>r.json());
  const nav=document.getElementById('gstatus-nav');
  _gmailMode=r.mode||'oauth';
  setTab(_gmailMode);
  if(r.configured){
    const lbl=r.mode==='oauth'?'OAuth':r.mode==='smtp'?'SMTP':'Browser';
    nav.textContent=`✅ ${r.user} (${lbl})`;nav.style.color='#16A34A';
    if(r.has_oauth && r.mode==='oauth') showOAuthConnected(r.user);
  } else {
    nav.textContent='⚙ Gmail not connected';nav.style.color='#888';
  }
  if(!r.oauth_configured){
    document.getElementById('oauth-not-configured').style.display='block';
    document.getElementById('redirect-uri-display').textContent=location.origin+'/auth/google/callback';
  }
}

/* ── Upload ── */
async function uploadExcel(el){
  const f=el.files[0];if(!f)return;
  const fd=new FormData();fd.append('file',f);
  const r=await fetch('/api/upload-excel',{method:'POST',body:fd}).then(r=>r.json());
  if(r.ok){toast('✅ Excel uploaded');loadRows();}else toast('Error: '+r.error,true);
  el.value='';
}
async function uploadTemplate(el){
  const f=el.files[0];if(!f)return;
  const fd=new FormData();fd.append('file',f);
  const r=await fetch('/api/upload-template',{method:'POST',body:fd}).then(r=>r.json());
  toast(r.ok?'✅ Template uploaded':'Error: '+r.error,!r.ok);el.value='';
}

/* ── Utilities ── */
async function post(url,data){
  const r=await fetch(url,{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify(data)});
  return r.json();
}
function toast(msg,err=false){
  const w=document.getElementById('toast-wrap');
  const d=document.createElement('div');d.className='toast'+(err?' err':'');d.textContent=msg;
  w.appendChild(d);setTimeout(()=>d.remove(),5000);
}
function prog(done,total,label){
  document.getElementById('pw').style.display='block';
  document.getElementById('pl').textContent=label;
  document.getElementById('pb').style.width=(total?Math.round(done/total*100):0)+'%';
}
function hideProg(){document.getElementById('pw').style.display='none';}
</script>
</body>
</html>"""

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
